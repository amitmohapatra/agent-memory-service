"""File ingestion: durable accept -> parse job -> nodes/chunks/edges -> raw archive."""

from __future__ import annotations

from pathlib import Path

import pytest

from memory_service.domain.context import MemoryExecutionContext
from memory_service.domain.enums import ArchiveStatus, ContextGraphEdge
from memory_service.domain.errors import DependencyUnavailable, ValidationFailed
from memory_service.modules.ingestion.service import IngestionService
from memory_service.modules.jobs.registry import register_handlers

pytestmark = pytest.mark.integration

FIXTURE = Path(__file__).resolve().parents[1] / "fixtures" / "acme_fy26_annual_report.md"
CTX = MemoryExecutionContext(tenant_id="acme", user_id="u1", workspace_id="ws1")


@pytest.fixture
def ingestion(container) -> IngestionService:
    register_handlers(container)
    return container.services["ingestion"]


async def test_accept_parse_and_archive(container, uow_factory, ingestion) -> None:
    async with uow_factory() as uow:
        ack = await ingestion.accept_file(
            uow,
            CTX,
            filename="acme_fy26_annual_report.md",
            media_type="text/markdown",
            data=FIXTURE.read_bytes(),
            title="ACME FY26",
        )
        await uow.commit()
    assert ack.job_ids and not ack.deduplicated
    async with uow_factory() as uow:
        doc = await uow.documents.get("acme", ack.document_id)
        assert doc is not None and doc.system_metadata["status"] == "STAGED"
        assert await uow.documents.staged_bytes("acme", ack.document_id) == FIXTURE.read_bytes()
    # run the parse job (recording queue -> drain)
    assert await container.tasks.drain() >= 1
    async with uow_factory() as uow:
        doc = await uow.documents.get("acme", ack.document_id)
        assert doc.system_metadata["status"] == "READY" and doc.current_version_id
        nodes = await uow.documents.list_nodes("acme", ack.document_id)
        chunks = await uow.documents.list_chunks("acme", ack.document_id)
        assert len(nodes) > 15 and len(chunks) >= 8
        page11 = next(c for c in chunks if c.page == 11 and "increased to EUR 98" in c.text)
        defined_by = await uow.documents.edges_from(
            "acme", [page11.node_id], kinds=[ContextGraphEdge.DEFINED_BY]
        )
        assert (
            defined_by
            and (await uow.documents.get_nodes("acme", [defined_by[0].target_id]))[0].page_start
            == 1
        )
        footnote = await uow.documents.edges_from(
            "acme", [page11.node_id], kinds=[ContextGraphEdge.FOOTNOTE]
        )
        assert (
            footnote
            and (await uow.documents.get_nodes("acme", [footnote[0].target_id]))[0].page_start == 20
        )
        # raw file archived + staged bytes purged
        assert doc.archive_status is ArchiveStatus.ARCHIVED
        assert await uow.documents.staged_bytes("acme", ack.document_id) is None
        seg = await uow.archive.get(f"seg_file_{ack.document_id}")
        assert seg is not None and seg.status == "VERIFIED" and seg.kind == "file"
        assert await container.blob.get(seg.bucket, seg.key) == FIXTURE.read_bytes()
    # the index job was chained (M6 attaches the handler)
    assert any(j.task_name == "document.index" for j in container.tasks.jobs.values())
    # authorization: owner reads, other user denied
    async with uow_factory() as uow:
        assert (
            await ingestion.get_document(uow, CTX, ack.document_id)
        ).document_id == ack.document_id
        from memory_service.domain.errors import ScopeDenied

        with pytest.raises(ScopeDenied):
            await ingestion.get_document(
                uow, MemoryExecutionContext(tenant_id="acme", user_id="u2"), ack.document_id
            )


async def test_duplicate_bytes_are_deduplicated_per_tenant(
    container, uow_factory, ingestion
) -> None:
    async with uow_factory() as uow:
        first = await ingestion.accept_file(
            uow, CTX, filename="a.md", media_type="text/markdown", data=b"# Same\n\ntext"
        )
        await uow.commit()
    async with uow_factory() as uow:
        second = await ingestion.accept_file(
            uow, CTX, filename="b.md", media_type="text/markdown", data=b"# Same\n\ntext"
        )
        other = await ingestion.accept_file(
            uow,
            MemoryExecutionContext(tenant_id="globex", user_id="u1"),
            filename="b.md",
            media_type="text/markdown",
            data=b"# Same\n\ntext",
        )
        await uow.commit()
    assert second.deduplicated and second.document_id == first.document_id
    assert not other.deduplicated and other.document_id != first.document_id


async def test_validation(container, uow_factory, ingestion) -> None:
    async with uow_factory() as uow:
        with pytest.raises(ValidationFailed):
            await ingestion.accept_file(
                uow, CTX, filename="x.bin", media_type="application/x-msdownload", data=b"MZ"
            )
        with pytest.raises(ValidationFailed):
            await ingestion.accept_file(
                uow, CTX, filename="x.md", media_type="text/markdown", data=b""
            )


async def test_blob_outage_keeps_staged_bytes_and_fails_job(
    container, uow_factory, ingestion
) -> None:
    async with uow_factory() as uow:
        ack = await ingestion.accept_file(
            uow, CTX, filename="n.md", media_type="text/markdown", data=b"# Note\n\nSome text."
        )
        await uow.commit()
    container.blob.available = False
    with pytest.raises(DependencyUnavailable):
        await ingestion.parse_document("acme", ack.document_id)
    async with uow_factory() as uow:
        doc = await uow.documents.get("acme", ack.document_id)
        assert doc.system_metadata["status"] == "READY"  # parsed content is durable
        assert doc.archive_status is ArchiveStatus.STAGED
        assert await uow.documents.staged_bytes("acme", ack.document_id) is not None
    container.blob.available = True
    assert (
        await ingestion.archive_raw_file("acme", ack.document_id) == f"seg_file_{ack.document_id}"
    )


@pytest.mark.skipif(
    not pytest.importorskip("docling", reason="docling not installed"), reason="docling"
)
async def test_docling_parses_docx(container, uow_factory) -> None:
    import docx  # python-docx (installed with docling)

    from memory_service.adapters.parsers.docling_parser import DoclingParser

    d = docx.Document()
    d.add_heading("Quarterly Update", level=1)
    d.add_heading("Revenue", level=2)
    d.add_paragraph("Revenue grew 9% to EUR 301 million. See Table 1.")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Segment", "FY26"
    table.cell(1, 0).text, table.cell(1, 1).text = "Subscriptions", "301"
    d.add_heading("Risks", level=2)
    d.add_paragraph("Churn increased in the Legacy Services unit.")
    import io

    buf = io.BytesIO()
    d.save(buf)
    parsed = await DoclingParser().parse(
        document_id="doc_x",
        tenant_id="acme",
        filename="update.docx",
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        data=buf.getvalue(),
    )
    titles = [n.title for n in parsed.nodes if n.title]
    assert "Revenue" in titles and "Risks" in titles
    assert any(
        n.representation.value == "TABLE" and "Subscriptions" in n.text for n in parsed.nodes
    )
    assert parsed.version.parser == "docling"
